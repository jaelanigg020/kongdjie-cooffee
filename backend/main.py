from fastapi import FastAPI, HTTPException, UploadFile, File, Form
from fastapi.middleware.cors import CORSMiddleware
from fastapi.staticfiles import StaticFiles
from pydantic import BaseModel
from typing import List, Optional
import sqlite3
import json
import os
import shutil
from datetime import datetime

app = FastAPI(title="Coffee Shop POS API")

app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)

from fastapi.responses import FileResponse, StreamingResponse
import io
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
FRONTEND_DIR = os.path.abspath(os.path.join(BASE_DIR, "..", "frontend"))
UPLOAD_DIR = os.path.join(BASE_DIR, "uploads")
DB_PATH = os.path.join(BASE_DIR, "pos.db")
os.makedirs(UPLOAD_DIR, exist_ok=True)

app.mount("/uploads", StaticFiles(directory=UPLOAD_DIR), name="uploads")
app.mount("/static", StaticFiles(directory=FRONTEND_DIR), name="static")

@app.get("/")
def serve_index():
    return FileResponse(os.path.join(FRONTEND_DIR, "index.html"))

@app.get("/style.css")
def serve_css():
    return FileResponse(os.path.join(FRONTEND_DIR, "style.css"))

@app.get("/app.js")
def serve_js():
    return FileResponse(os.path.join(FRONTEND_DIR, "app.js"))

@app.get("/logo.png")
def serve_logo():
    return FileResponse(os.path.join(FRONTEND_DIR, "logo.png"))

def get_db():
    conn = sqlite3.connect(DB_PATH)
    conn.row_factory = sqlite3.Row
    return conn

def init_db():
    conn = get_db()
    c = conn.cursor()
    
    # Products
    c.execute("""
        CREATE TABLE IF NOT EXISTS products (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            name TEXT NOT NULL,
            category TEXT NOT NULL,
            price REAL NOT NULL,
            image_url TEXT,
            stock INTEGER DEFAULT 100,
            is_active INTEGER DEFAULT 1
        )
    """)
    
    # Shifts
    c.execute("""
        CREATE TABLE IF NOT EXISTS shifts (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            cashier_name TEXT NOT NULL,
            start_time TEXT NOT NULL,
            end_time TEXT,
            starting_cash REAL NOT NULL,
            ending_cash REAL,
            actual_cash REAL,
            status TEXT DEFAULT 'OPEN'
        )
    """)
    
    # Orders
    c.execute("""
        CREATE TABLE IF NOT EXISTS orders (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            order_number TEXT UNIQUE NOT NULL,
            shift_id INTEGER,
            order_type TEXT NOT NULL, -- Dine In, Take Away, Delivery
            table_number TEXT,
            total_amount REAL NOT NULL,
            tax REAL DEFAULT 0,
            discount REAL DEFAULT 0,
            payment_method TEXT, -- Cash, EDC BCA (Card/QRIS), Transfer
            payment_status TEXT DEFAULT 'UNPAID', -- UNPAID, PAID, CANCELLED
            notes TEXT,
            created_at TEXT NOT NULL
        )
    """)
    
    # Order Items
    c.execute("""
        CREATE TABLE IF NOT EXISTS order_items (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            order_id INTEGER NOT NULL,
            product_id INTEGER NOT NULL,
            product_name TEXT NOT NULL,
            price REAL NOT NULL,
            quantity INTEGER NOT NULL,
            notes TEXT,
            FOREIGN KEY (order_id) REFERENCES orders (id)
        )
    """)
    
    # Seed initial products if empty
    c.execute("SELECT COUNT(*) FROM products")
    if c.fetchone()[0] == 0:
        sample_products = [
            ("Espresso Single", "Coffee", 18000, "https://images.unsplash.com/photo-1510591509098-f4fdc6d0ff04?w=300", 50),
            ("Americano / Long Black", "Coffee", 22000, "https://images.unsplash.com/photo-1551030173-122aabc4489c?w=300", 100),
            ("Cafe Latte", "Coffee", 28000, "https://images.unsplash.com/photo-1570968915860-54d5c301fa9f?w=300", 80),
            ("Caramel Macchiato", "Coffee", 32000, "https://images.unsplash.com/photo-1485808191679-5f86510681a2?w=300", 60),
            ("Matcha Latte", "Non-Coffee", 30000, "https://images.unsplash.com/photo-1536256263959-770b48d82b0a?w=300", 40),
            ("Earl Grey Tea", "Non-Coffee", 20000, "https://images.unsplash.com/photo-1576092768241-dec231879fc3?w=300", 90),
            ("Butter Croissant", "Pastry", 25000, "https://images.unsplash.com/photo-1555507036-ab1f4038808a?w=300", 25),
            ("Chocolate Brownie", "Pastry", 22000, "https://images.unsplash.com/photo-1606313564200-e75d5e30476c?w=300", 30)
        ]
        c.executemany("INSERT INTO products (name, category, price, image_url, stock) VALUES (?, ?, ?, ?, ?)", sample_products)
        
    conn.commit()
    conn.close()

init_db()

# Models
class ProductCreate(BaseModel):
    name: str
    category: str
    price: float
    image_url: Optional[str] = None
    stock: int = 100

class ProductUpdate(BaseModel):
    name: Optional[str] = None
    category: Optional[str] = None
    price: Optional[float] = None
    image_url: Optional[str] = None
    stock: Optional[int] = None
    is_active: Optional[int] = None

class OrderItemSchema(BaseModel):
    product_id: int
    product_name: str
    price: float
    quantity: int
    notes: Optional[str] = ""

class OrderCreate(BaseModel):
    order_type: str
    table_number: Optional[str] = "-"
    items: List[OrderItemSchema]
    tax: float = 0
    discount: float = 0
    payment_method: Optional[str] = None
    payment_status: str = "PAID"
    notes: Optional[str] = ""
    shift_id: Optional[int] = 1

class ShiftStart(BaseModel):
    cashier_name: str
    starting_cash: float

class ShiftClose(BaseModel):
    shift_id: int
    actual_cash: float

# Routes
@app.get("/api/products")
def get_products(category: Optional[str] = None):
    conn = get_db()
    c = conn.cursor()
    if category:
        c.execute("SELECT * FROM products WHERE category = ? AND is_active = 1", (category,))
    else:
        c.execute("SELECT * FROM products WHERE is_active = 1")
    rows = [dict(r) for r in c.fetchall()]
    conn.close()
    return rows

@app.post("/api/products")
def create_product(product: ProductCreate):
    conn = get_db()
    c = conn.cursor()
    c.execute(
        "INSERT INTO products (name, category, price, image_url, stock) VALUES (?, ?, ?, ?, ?)",
        (product.name, product.category, product.price, product.image_url or "", product.stock)
    )
    product_id = c.lastrowid
    conn.commit()
    conn.close()
    return {"id": product_id, **product.dict()}

@app.put("/api/products/{product_id}")
def update_product(product_id: int, product: ProductUpdate):
    conn = get_db()
    c = conn.cursor()
    fields = []
    values = []
    for k, v in product.dict(exclude_unset=True).items():
        fields.append(f"{k} = ?")
        values.append(v)
    if not fields:
        raise HTTPException(status_code=400, detail="No fields provided")
    values.append(product_id)
    c.execute(f"UPDATE products SET {', '.join(fields)} WHERE id = ?", values)
    conn.commit()
    conn.close()
    return {"status": "success", "id": product_id}

@app.delete("/api/products/{product_id}")
def delete_product(product_id: int):
    conn = get_db()
    c = conn.cursor()
    c.execute("UPDATE products SET is_active = 0 WHERE id = ?", (product_id,))
    conn.commit()
    conn.close()
    return {"status": "deleted"}

@app.post("/api/upload")
async def upload_file(file: UploadFile = File(...)):
    filename = f"{datetime.now().strftime('%Y%m%d%H%M%S')}_{file.filename}"
    file_path = os.path.join(UPLOAD_DIR, filename)
    with open(file_path, "wb") as buffer:
        shutil.copyfileobj(file.file, buffer)
    return {"url": f"http://127.0.0.1:8000/uploads/{filename}"}

@app.post("/api/orders")
def create_order(order: OrderCreate):
    conn = get_db()
    c = conn.cursor()
    
    order_num = f"ORD-{datetime.now().strftime('%y%m%d%H%M%S')}"
    subtotal = sum(item.price * item.quantity for item in order.items)
    total_amount = subtotal + order.tax - order.discount
    created_at = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    
    c.execute("""
        INSERT INTO orders (order_number, shift_id, order_type, table_number, total_amount, tax, discount, payment_method, payment_status, notes, created_at)
        VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
    """, (
        order_num, order.shift_id, order.order_type, order.table_number,
        total_amount, order.tax, order.discount, order.payment_method,
        order.payment_status, order.notes, created_at
    ))
    order_id = c.lastrowid
    
    for item in order.items:
        c.execute("""
            INSERT INTO order_items (order_id, product_id, product_name, price, quantity, notes)
            VALUES (?, ?, ?, ?, ?, ?)
        """, (order_id, item.product_id, item.product_name, item.price, item.quantity, item.notes))
        
        # Deduct stock
        c.execute("UPDATE products SET stock = MAX(0, stock - ?) WHERE id = ?", (item.quantity, item.product_id))
        
    conn.commit()
    conn.close()
    
    return {
        "id": order_id,
        "order_number": order_num,
        "total_amount": total_amount,
        "created_at": created_at,
        "status": "success"
    }

@app.get("/api/orders")
def get_orders(limit: int = 50):
    conn = get_db()
    c = conn.cursor()
    c.execute("SELECT * FROM orders ORDER BY id DESC LIMIT ?", (limit,))
    orders = [dict(r) for r in c.fetchall()]
    
    for ord in orders:
        c.execute("SELECT * FROM order_items WHERE order_id = ?", (ord["id"],))
        ord["items"] = [dict(r) for r in c.fetchall()]
        
    conn.close()
    return orders

@app.get("/api/reports/export-excel")
def export_orders_excel():
    conn = get_db()
    c = conn.cursor()
    c.execute("SELECT * FROM orders ORDER BY id DESC")
    orders = [dict(r) for r in c.fetchall()]
    
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Riwayat Transaksi"
    
    # Header Styling
    header_fill = PatternFill(start_color="191C1F", end_color="191C1F", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    thin_border = Border(
        left=Side(style='thin', color='E0E0E0'),
        right=Side(style='thin', color='E0E0E0'),
        top=Side(style='thin', color='E0E0E0'),
        bottom=Side(style='thin', color='E0E0E0')
    )
    
    headers = [
        "No. Order", "Waktu Transaksi", "Tipe Pesanan", "No. Meja",
        "Detail Item & Qty", "Subtotal (Rp)", "Pajak PB1 (Rp)", "Total Bayar (Rp)",
        "Metode Pembayaran", "Status"
    ]
    
    ws.append(headers)
    for col_idx in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    
    # Rows
    for ord in orders:
        c.execute("SELECT product_name, quantity, price FROM order_items WHERE order_id = ?", (ord["id"],))
        items = c.fetchall()
        items_str = ", ".join([f"{item['product_name']} x{item['quantity']}" for item in items])
        subtotal = ord["total_amount"] - ord["tax"] + ord["discount"]
        
        row = [
            ord["order_number"],
            ord["created_at"],
            ord["order_type"],
            ord["table_number"] or "-",
            items_str,
            subtotal,
            ord["tax"],
            ord["total_amount"],
            ord["payment_method"],
            ord["payment_status"]
        ]
        ws.append(row)
        
    for row in ws.iter_rows(min_row=2, max_row=len(orders) + 1, min_col=1, max_col=len(headers)):
        for cell in row:
            cell.border = thin_border
            if cell.column in [6, 7, 8]:
                cell.number_format = '#,##0'
                
    # Auto adjust column width
    for col in ws.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = openpyxl.utils.get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = max(max_len + 4, 12)
        
    conn.close()
    
    stream = io.BytesIO()
    wb.save(stream)
    stream.seek(0)
    
    filename = f"Laporan_Transaksi_KongDjie_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
    return StreamingResponse(
        stream,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename={filename}"}
    )

@app.get("/api/reports/daily")
def get_daily_report():
    conn = get_db()
    c = conn.cursor()
    today = datetime.now().strftime("%Y-%m-%d")
    
    c.execute("""
        SELECT 
            COUNT(*) as total_orders,
            COALESCE(SUM(total_amount), 0) as total_revenue,
            COALESCE(SUM(CASE WHEN payment_method = 'Cash' THEN total_amount ELSE 0 END), 0) as cash_total,
            COALESCE(SUM(CASE WHEN payment_method LIKE '%EDC%' THEN total_amount ELSE 0 END), 0) as edc_total
        FROM orders
        WHERE created_at LIKE ? AND payment_status = 'PAID'
    """, (f"{today}%",))
    summary = dict(c.fetchone())
    
    c.execute("""
        SELECT oi.product_name, SUM(oi.quantity) as qty, SUM(oi.price * oi.quantity) as sales
        FROM order_items oi
        JOIN orders o ON oi.order_id = o.id
        WHERE o.created_at LIKE ? AND o.payment_status = 'PAID'
        GROUP BY oi.product_name
        ORDER BY qty DESC
    """, (f"{today}%",))
    top_products = [dict(r) for r in c.fetchall()]
    
    conn.close()
    return {"summary": summary, "top_products": top_products}
